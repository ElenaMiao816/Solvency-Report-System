from __future__ import annotations

import io
import unittest

import openpyxl

from services.solvency_ai_table_extractor import (
    AIExtractionBundle,
    PageGrid,
    extract_unit_records,
    reconstructed_workbook_bytes,
)
from services.solvency_table_extractor import ExtractedTable


class UnitExtractionTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            [
                "\u884c\u6b21",
                "\u9879\u76ee",
                "\u671f\u672b\u6570\uff08\u5143\uff09",
                "\u671f\u521d\u6570",
            ],
            ["1", "\u6838\u5fc3\u4e00\u7ea7\u8d44\u672c\uff08\u5143\uff09", "100", "90"],
            ["2", "\u6838\u5fc3\u4e8c\u7ea7\u8d44\u672c", "80", "70"],
            [
                "3",
                "\u6838\u5fc3\u507f\u4ed8\u80fd\u529b\u5145\u8db3\u7387",
                "130%",
                "120%",
            ],
        ]
        self.grid = PageGrid(
            12,
            "\n".join([
                "000|\u5b9e\u9645\u8d44\u672c\u8868",
                "001|\u5355\u4f4d\uff1a\u4eba\u6c11\u5e01\u5143",
                "002|\u884c\u6b21 \u9879\u76ee \u671f\u672b\u6570 \u671f\u521d\u6570",
                "003|\u6838\u5fc3\u4e00\u7ea7\u8d44\u672c 100 90",
                "004|\u5b9e\u9645\u8d44\u672c\u5408\u8ba1 180 160",
            ]),
            30,
        )

    def _table(self) -> ExtractedTable:
        units = extract_unit_records("ACTUAL_CAPITAL", self.rows, [self.grid])
        return ExtractedTable(
            "ACTUAL_CAPITAL",
            "\u5b9e\u9645\u8d44\u672c\u8868",
            12,
            1,
            self.rows,
            source_pages=[12],
            unit_records=units,
        )

    def test_extracts_table_column_row_and_percent_units(self):
        units = self._table().unit_records
        found = {(item.scope, item.normalized_unit) for item in units}
        self.assertIn(("\u8868\u7ea7", "\u5143"), found)
        self.assertIn(("\u5217\u7ea7", "\u5143"), found)
        self.assertIn(("\u884c\u7ea7", "\u5143"), found)
        self.assertTrue(
            any(
                item.normalized_unit == "%"
                and item.target == "\u6838\u5fc3\u507f\u4ed8\u80fd\u529b\u5145\u8db3\u7387"
                for item in units
            )
        )

    def test_unit_footer_does_not_modify_data_rows(self):
        table = self._table()
        self.assertEqual(len(table.to_frame()), len(self.rows))
        preview = table.to_frame(include_unit_footer=True)
        self.assertEqual(len(preview), len(self.rows) + 1)
        self.assertEqual(preview.iloc[-1, 0], "\u3010\u5355\u4f4d\u5907\u6ce8\u3011")

    def test_workbook_contains_unit_sheet_and_footer(self):
        table = self._table()
        payload = reconstructed_workbook_bytes(AIExtractionBundle([table], []))
        workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
        self.assertIn("\u5355\u4f4d\u4fe1\u606f", workbook.sheetnames)
        table_sheet = workbook["\u5b9e\u9645\u8d44\u672c\u8868_P12"]
        self.assertEqual(
            table_sheet.cell(table_sheet.max_row, 1).value,
            "\u3010\u5355\u4f4d\u5907\u6ce8\u3011",
        )


if __name__ == "__main__":
    unittest.main()

